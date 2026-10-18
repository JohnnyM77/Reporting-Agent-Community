"""Build outputs/archetype_review.xlsx from the classifier and email it.

Syncs config/ticker_archetypes.yaml first (new tickers are classified and
written), classifies any remaining active ticker so every row is fresh, then
writes the workbook, lowest confidence first, with a blank "my call" column.

    python scripts/archetype_review.py            # build + email
    python scripts/archetype_review.py --no-email
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from shared import archetype_sync, classifier  # noqa: E402

OUT = Path("outputs/archetype_review.xlsx")
SUBJECT = "Archetype review: please confirm"
COLUMNS = ["ticker", "company", "watchlists", "suggested archetype", "base multiple", "modifiers applied",
           "final buy multiple", "10-year median P/E", "current P/E", "current price", "implied buy price",
           "confidence", "my call", "notes"]


def build_rows(entries: dict, results: dict, rules: dict) -> list[dict]:
    rows = []
    for t, e in entries.items():
        if not e.get("active", True):
            continue
        c = results[t]
        f = c.features or classifier.Features(ticker=t)
        mult = c.buy_multiple
        unclassified = c.archetype == classifier.UNCLASSIFIED
        rows.append({
            "ticker": t,
            "company": f.company or e.get("company", ""),
            "watchlists": ", ".join(e.get("lists") or []),
            "suggested archetype": "UNCLASSIFIED" if unclassified else classifier.archetype_label(rules, c.archetype),
            "base multiple": c.base_multiple,
            "modifiers applied": c.modifiers_text() if c.base_multiple is not None else "",
            "final buy multiple": mult if mult is not None else (f"no P/E line ({c.alt_metric})" if c.alt_metric else None),
            "10-year median P/E": round(c.median_pe_10y, 1) if c.median_pe_10y else None,
            "current P/E": round(f.current_pe, 1) if f.current_pe else None,
            "current price": f.price,
            "implied buy price": c.implied_buy_price,
            "confidence": c.confidence,
            "my call": "",
            "notes": " | ".join(x for x in (c.reasons + c.guardrails) if x),
            "_arch": c.archetype,
        })
    rows.sort(key=lambda r: (r["confidence"], r["ticker"]))
    return rows


def write_xlsx(rows: list[dict], rules: dict, path: Path = OUT) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2D4E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [10, 30, 28, 22, 9, 46, 14, 12, 10, 10, 12, 11, 22, 90]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    low = PatternFill("solid", fgColor="FDE9D9")
    for i, r in enumerate(rows, 2):
        ws.cell(i, 12).number_format = "0%"
        if r["confidence"] < 0.6:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(i, col).fill = low
    names = ",".join(list((rules.get("archetypes") or {})))
    dv = DataValidation(type="list", formula1=f'"{names}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{len(rows) + 1}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-email", action="store_true")
    args = ap.parse_args()
    rules = classifier.load_rules()
    results: dict = {}
    out = archetype_sync.sync_archetypes(results=results)
    entries = out["entries"]
    todo = [t for t, e in entries.items() if e.get("active", True) and t not in results]
    print(f"[review] re-classifying {len(todo)} existing ticker(s)", flush=True)

    def run(t):
        try:
            return classifier.classify_ticker(t, rules, entry=entries.get(t))
        except Exception as exc:  # noqa: BLE001
            return archetype_sync.failed_classification(t, exc)
    with ThreadPoolExecutor(max_workers=4) as pool:
        for t, c in zip(todo, pool.map(run, todo)):
            results[t] = c

    rows = build_rows(entries, results, rules)
    write_xlsx(rows, rules)
    counts = Counter(r["_arch"] for r in rows)
    print(f"[review] {len(rows)} tickers -> {OUT}")
    print("[review] by archetype: " + ", ".join(f"{k}={v}" for k, v in counts.most_common()))
    print("[review] 20 lowest confidence:")
    for r in rows[:20]:
        print(f"  {r['ticker']:<10} {r['confidence']:.2f}  {r['suggested archetype']:<24} {r['notes'][:110]}")

    if not args.no_email:
        from shared.email_service import send_email
        body = (f"{len(rows)} tickers classified. Lowest confidence first; fill in 'my call' and send back.\n\n"
                + "\n".join(f"{k}: {v}" for k, v in counts.most_common()))
        send_email(SUBJECT, body, "<pre>" + body + "</pre>", attachments=[OUT], raise_on_error=True)
        print("[review] email sent", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
