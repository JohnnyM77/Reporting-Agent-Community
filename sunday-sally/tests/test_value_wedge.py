import datetime as dt
from pathlib import Path

import pandas as pd
import pytest

from src.alert_engine import classify_alert
from src.pathing import repo_root
from src.value_wedge import (
    EarningsPoint, MultipleRegime, WedgeInputs, compute_wedge, load_wedge_inputs, normalise_ticker, zone_for,
)
from src.wedge_report import wedge_html

D = dt.date


def _prices(start="2020-01-01", end="2023-12-29", value=10.0, last=None):
    idx = pd.bdate_range(start, end)
    s = pd.Series(value, index=idx, dtype=float)
    if last is not None:
        s.iloc[-1] = last
    return s


def _inputs(earnings, multiples, prices, notes=None, source="config"):
    return WedgeInputs("TST.AX", "Test Co", "A", "reported_ttm", earnings, multiples, prices,
                       notes or [], multiples_source=source)


def test_eps_growth_moves_corridor_but_not_relative_wedge():
    """EPS 20c -> 40c at constant 25x/45x: dollar wedge doubles, relative wedge stays 80%."""
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0), EarningsPoint(D(2022, 1, 3), "FY21", 40.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    r = compute_wedge(_inputs(e, m, _prices()))
    early = next(p for p in r.points if p.date.year == 2021)
    late = r.points[-1]
    assert early.buy_value == pytest.approx(5.0) and early.sell_value == pytest.approx(9.0)
    assert late.buy_value == pytest.approx(10.0) and late.sell_value == pytest.approx(18.0)
    assert late.abs_wedge == pytest.approx(2 * early.abs_wedge)
    assert early.rel_wedge == late.rel_wedge == pytest.approx(0.8)
    ev = [x for x in r.events if x.kind == "earnings"]
    assert len(ev) == 1 and "shifted upward" in ev[0].text and "relative wedge unchanged" in ev[0].text


def test_multiple_change_widens_wedge_with_eps_flat():
    """EPS flat at 30c; 20x/30x -> 25x/45x: a framework change, relative wedge 50% -> 80%."""
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 30.0)]
    m = [MultipleRegime(D(2020, 1, 1), 20, 30),
         MultipleRegime(D(2022, 6, 1), 25, 45, "A&D now the largest division")]
    r = compute_wedge(_inputs(e, m, _prices(), source="config multiple_history"))
    before = next(p for p in r.points if p.date.year == 2021)
    after = r.points[-1]
    assert (before.buy_value, before.sell_value) == (pytest.approx(6.0), pytest.approx(9.0))
    assert (after.buy_value, after.sell_value) == (pytest.approx(7.5), pytest.approx(13.5))
    assert before.rel_wedge == pytest.approx(0.5) and after.rel_wedge == pytest.approx(0.8)
    mv = [x for x in r.events if x.kind == "multiple"]
    assert len(mv) == 1
    assert "SELL multiple increased from 30x to 45x" in mv[0].text
    assert "corridor widened" in mv[0].text and "not an earnings change" in mv[0].text
    assert "A&D" in mv[0].text
    widened = dict(r.diagnostics)["Has the valuation range itself widened or narrowed?"]
    assert widened.startswith("Widened")


def test_no_look_ahead_eps_applies_from_release_date_only():
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0), EarningsPoint(D(2021, 8, 20), "FY21", 40.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    r = compute_wedge(_inputs(e, m, _prices()))
    for p in r.points:
        assert p.eps_used == (40.0 if p.date >= D(2021, 8, 20) else 20.0)


def test_points_before_first_eps_have_no_corridor():
    e = [EarningsPoint(D(2021, 1, 4), "FY20", 20.0)]
    m = [MultipleRegime(D(2021, 1, 4), 25, 45)]
    r = compute_wedge(_inputs(e, m, _prices()))
    assert r.points[0].buy_value is None and r.points[0].zone == "No valuation"


def test_zones_and_above_sell_diagnostics():
    assert zone_for(-0.1) == "Below BUY"
    assert zone_for(0.5) == "Middle of wedge"
    assert zone_for(0.9) == "Upper third of wedge"
    assert zone_for(1.2) == "Above SELL"
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    # Steady $8 then a jump to $12 on the last day: a new 52w high AND above SELL ($9).
    r = compute_wedge(_inputs(e, m, _prices(value=8.0, last=12.0)))
    cur = r.current
    assert cur.zone == "Above SELL" and cur.actual_pe == pytest.approx(60)
    assert r.near_52w_high
    a7 = r.diagnostics[-1][1]
    assert "genuine valuation" in a7
    assert dict(r.diagnostics)["Has it crossed the SELL boundary?"].startswith("Yes")


def test_52w_high_inside_wedge_is_not_a_valuation_sell():
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    r = compute_wedge(_inputs(e, m, _prices(value=5.5, last=6.0)))   # $6 = 30x, lower third
    assert r.near_52w_high and r.current.zone == "Lower third of wedge"
    assert "price event, not a valuation sell" in r.diagnostics[-1][1]


def test_alert_above_sell_triggers_without_52w_high():
    w = {"position": 1.4, "actual_pe": 60.0, "sell_pe": 45.0, "sell_value": 9.0}
    a = classify_alert(0.20, 0.05, None, None, None, None, 0.6, wedge=w)
    assert a.triggered and a.tier == "Tier 2: Review"
    assert "above SELL" in a.reasons[0] and "not near the 52-week high" in a.reasons[0]
    # inside the wedge and far from the high: nothing
    w_in = dict(w, position=0.5)
    assert not classify_alert(0.20, 0.05, None, None, None, None, 0.6, wedge=w_in).triggered
    # near the high and above SELL: the wedge counts as a stretch signal
    near = classify_alert(0.01, 0.05, None, None, None, None, 0.6, wedge=w)
    assert near.tier == "Tier 2: Review" and any("above SELL" in r for r in near.reasons)
    # unchanged behaviour without a wedge
    assert classify_alert(0.01, 0.05, None, None, None, None, 0.6).tier == "Tier 1: Watch"


def test_pwh_config_loads_with_unchanged_framework():
    px = _prices("2017-01-02", "2026-09-30", value=12.0)
    inp = load_wedge_inputs("PWH", repo=repo_root(), prices=px)
    assert inp.ticker == "PWH.AX"
    assert [(m.buy_pe, m.sell_pe) for m in inp.multiples] == [(25.0, 50.0)]   # JM's framework, Oct 2026
    assert inp.fair_split == 0.35 and not inp.smoothing                       # statutory EPS, no smoothing
    assert inp.depressed and inp.depressed[0]["reference_period"] == "FY2024"
    assert any("Brisbane relocation" in w and "BUY $6.17" in w for w in inp.warnings)
    assert inp.eps_basis == "reported_ttm"
    # Reported EPS, release-dated, ascending; FY25 reset is in there.
    fy25 = next(e for e in inp.earnings if e.period == "FY2025")
    assert fy25.eps_cents == pytest.approx(9.712) and fy25.date == D(2025, 8, 25)
    assert not any(e.notes.lower().startswith(("est", "derived")) for e in inp.earnings)
    # Every annual figure comes from the SWS export; the one half-year TTM is
    # computed from two reported halves, and the missing halves are listed.
    assert all(e.source.startswith("SWS") for e in inp.earnings)
    ttm = [e for e in inp.earnings if "TTM" in e.period]
    assert [(e.period, e.eps_cents) for e in ttm] == [("1H2026 TTM", pytest.approx(9.712 - 4.1 + 5.6))]
    assert any("1H2016-1H2024" in w for w in inp.warnings)
    r = compute_wedge(inp)
    assert r.current.rel_wedge == pytest.approx(1.0)
    assert r.current.fair_pe == pytest.approx(33.75)


def test_normalise_ticker():
    assert normalise_ticker("pwh") == "PWH.AX"
    assert normalise_ticker("ASX:PWH") == "PWH.AX"
    assert normalise_ticker("RR.L") == "RR.L"


def test_wedge_html_is_email_safe():
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    html = wedge_html(compute_wedge(_inputs(e, m, _prices(value=8.0, last=12.0))), cid="wedge_tst")
    for banned in ("<style", "class=", "display:flex", "display:grid"):
        assert banned not in html
    assert "cid:wedge_tst" in html and "Above SELL" in html


def test_chart_renders(tmp_path: Path):
    pytest.importorskip("matplotlib")
    from src.wedge_chart import render_wedge_chart
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0), EarningsPoint(D(2022, 1, 3), "FY21", 40.0)]
    m = [MultipleRegime(D(2020, 1, 1), 20, 30), MultipleRegime(D(2022, 6, 1), 25, 45, "review")]
    out = render_wedge_chart(compute_wedge(_inputs(e, m, _prices())), tmp_path / "w.png")
    assert out and out.stat().st_size > 10_000


def test_half_year_ttm_needs_both_reported_halves():
    from src.value_wedge import half_year_ttm
    annual = [EarningsPoint(D(2024, 8, 20), "FY2024", 20.0), EarningsPoint(D(2025, 8, 20), "FY2025", 10.0)]
    cfg = {"half_year_eps": [
        {"period": "1H2025", "release": "2025-02-20", "eps": 4.0},
        {"period": "1H2026", "release": "2026-02-20", "eps": 6.0},
    ]}
    pts, missing = half_year_ttm(cfg, annual)
    # 1H2025 has no 1H2024 beside it, so no TTM there; 1H2026 = 10 - 4 + 6.
    assert [(p.period, p.eps_cents, p.date) for p in pts] == [("1H2026 TTM", 12.0, D(2026, 2, 20))]
    assert "1H2024" in missing and "1H2025" not in missing


def test_sws_import_reads_unix_dates_and_annual_eps(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from src.sws_import import annual_eps_rows, read_sws_rows, replace_earnings_block, render_earnings_block, unix_to_date
    assert unix_to_date(1782777600) == D(2026, 6, 30)
    wb = openpyxl.Workbook()
    ws = wb.active
    data = {
        "past_last_earnings_update_annual": [1782777600],                    # 2026-06-30
        "past_earnings_per_share_ltm_history_offset": [0, -1, -2],
        "past_earnings_per_share_ltm_history": [0.17792, 0.09712, 0.24693],
        "dividend_historical_dividend_payments_date": [1756080000, 1787184000],  # 2025-08-25, 2026-08-20
        "dividend_payout_ratio_history_date": [1782777600],
        "dividend_payout_ratio_history": [0.44965],
    }
    for r, (k, vals) in enumerate(data.items(), 1):
        ws.cell(r, 1, k)
        for c, v in enumerate(vals, 2):
            ws.cell(r, c, v)
    path = tmp_path / "261003_SWS_Data.xlsx"
    wb.save(path)
    rows = annual_eps_rows(read_sws_rows(path), "SWS export 261003")
    assert [r["period"] for r in rows] == ["FY2024", "FY2025", "FY2026"]
    fy26, fy25, fy24 = rows[2], rows[1], rows[0]
    assert fy26["ttm_eps"] == pytest.approx(17.792) and fy26["date"] == "2026-08-20"
    assert fy26["ttm_div"] == pytest.approx(8.0, abs=0.01)
    assert fy25["date"] == "2025-08-25"
    # No SWS date for FY24: the ASX two-month deadline, never an earlier guess.
    assert fy24["date"] == "2024-08-31" and "deadline" in fy24["notes"]
    yaml_text = "ticker: X\nearnings:\n  - {date: old}\n\nhalf_year_eps:\n  - {period: 1H2026}\n"
    new = replace_earnings_block(yaml_text, render_earnings_block(rows))
    assert "old" not in new and "half_year_eps:\n  - {period: 1H2026}" in new and new.count("FY2026") == 1


def test_workbook_is_formula_driven_with_native_charts(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from src.wedge_workbook import build_wedge_workbook
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0), EarningsPoint(D(2022, 1, 3), "FY21", 40.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    out = build_wedge_workbook(compute_wedge(_inputs(e, m, _prices())), tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Read me", "Chart", "Weekly", "EPS", "Multiples", "Story"]
    ws = wb["Weekly"]
    assert ws["B2"].value == 10.0                               # price is a value
    assert str(ws["C2"].value).startswith("=IFERROR(LOOKUP(2,1/")  # EPS in force is a formula
    assert ws["F2"].value.startswith("=IF(") and "D2" in ws["F2"].value
    assert wb["Multiples"]["B2"].value == 25 and wb["EPS"]["C3"].value == 40.0
    assert len(wb["Chart"]._charts) == 4


def test_chart_also_writes_editable_svg(tmp_path: Path):
    pytest.importorskip("matplotlib")
    from src.wedge_chart import render_wedge_chart
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 45)]
    render_wedge_chart(compute_wedge(_inputs(e, m, _prices())), tmp_path / "w.png")
    svg = (tmp_path / "w.svg").read_text()
    assert "<text" in svg and "the wedge between BUY and SELL" in svg   # real text, not outlines


def test_foreign_currency_eps_is_converted_at_release_and_stays_flat():
    """US$ EPS, A$ price: converted once at the release-date rate, so BUY/SELL
    are flat until the next result even though the exchange rate moves."""
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 100.0),                 # 100 US cents
         EarningsPoint(D(2022, 3, 1), "1H22 TTM", 100.0)]
    m = [MultipleRegime(D(2020, 1, 1), 10, 20)]
    px = _prices(value=20.0)
    fx = pd.Series(1.5, index=px.index)
    fx[fx.index >= "2022-01-01"] = 1.4
    inp = WedgeInputs("TST.AX", "Test", "A", "reported_ttm", e, m, px,
                      eps_currency="USD", fx=fx, fx_label="AUD per USD")
    r = compute_wedge(inp)
    early = next(p for p in r.points if p.date.year == 2021)
    jan22 = next(p for p in r.points if p.date.year == 2022 and p.date.month == 1)
    late = r.points[-1]
    assert early.reported_eps == 100.0 and early.eps_used == pytest.approx(150.0)
    assert early.buy_value == pytest.approx(15.0) and early.sell_value == pytest.approx(30.0)
    assert jan22.fx == pytest.approx(1.5) and jan22.buy_value == pytest.approx(15.0)   # FX moved, BUY did not
    assert late.buy_value == pytest.approx(14.0) and late.fx == pytest.approx(1.4)     # next result, new rate
    from src.value_wedge import daily_corridor
    day = daily_corridor(r)
    held = day[(day.index >= "2020-01-01") & (day.index < "2022-03-01")]
    assert held["buy"].nunique() == 1 and held["sell"].nunique() == 1
    assert "exchange rate" in r.diagnostics[0][1] or "Flat" in r.diagnostics[0][1]
    with pytest.raises(ValueError):
        compute_wedge(WedgeInputs("TST.AX", "Test", "A", "reported_ttm", e, m, px, eps_currency="USD"))


def test_loss_years_remove_the_corridor_and_say_so():
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0), EarningsPoint(D(2021, 1, 4), "FY20", -10.0),
         EarningsPoint(D(2022, 1, 3), "FY21", 30.0)]
    m = [MultipleRegime(D(2020, 1, 1), 10, 20)]
    r = compute_wedge(_inputs(e, m, _prices()))
    loss = next(p for p in r.points if p.date.year == 2021 and p.date.month == 6)
    assert loss.buy_value is None and loss.zone == "No valuation" and loss.actual_pe is None
    texts = [x.text for x in r.events if x.kind == "earnings"]
    assert any("turned to a loss" in t for t in texts) and any("back to a profit" in t for t in texts)


def test_start_year_before_first_buy_from_the_ledger(tmp_path):
    import json
    from src.value_wedge import decision_notes, resolve_start
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "decisions.json").write_text(json.dumps({"decisions": [
        {"ticker": "TST", "date": "2021-11-04"}, {"ticker": "TST", "date": "2021-08-14"},
        {"ticker": "OTHER", "date": "2019-01-01"}]}))
    notes = decision_notes(tmp_path, "TST.AX")
    assert [n.date for n in notes] == [D(2021, 8, 14), D(2021, 11, 4)]
    assert notes[0].label == "Bought (decision 1 of 2)" and notes[0].kind == "decision"
    assert resolve_start("year_before_buy", notes) == D(2020, 8, 14)
    assert resolve_start("full", notes) is None and resolve_start("2019-05-01", notes) == D(2019, 5, 1)
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 10, 20)]
    inp = _inputs(e, m, _prices(), notes=notes)
    inp.start = D(2020, 8, 14)
    r = compute_wedge(inp)
    assert r.points[0].date >= D(2020, 8, 14)
    assert [x.kind for x in r.events].count("decision") == 2


def test_coverage_gaps_list_missing_halves_and_years():
    from src.value_wedge import eps_coverage_gaps
    e = [EarningsPoint(D(2014, 8, 19), "FY2014", 260.0), EarningsPoint(D(2016, 8, 16), "FY2016", -120.0),
         EarningsPoint(D(2017, 2, 21), "1H2017 TTM", 46.7), EarningsPoint(D(2018, 2, 20), "1H2018 TTM", 94.0)]
    g = eps_coverage_gaps(e)
    assert g["FY"] == ["FY2015", "FY2017"] and g["1H"] == ["1H2015", "1H2016"]


def test_bhp_config_is_sourced_usd_eps():
    px = _prices("2013-08-01", "2026-09-30", value=45.0)
    fx = pd.Series(1.5, index=px.index)
    inp = load_wedge_inputs("BHP", repo=repo_root(), prices=px, fx=fx, start="year_before_buy")
    assert inp.eps_currency == "USD" and inp.start == D(2014, 8, 14)        # first buy 14 Aug 2015
    assert all(e.source for e in inp.earnings)
    assert not any(e.notes.lower().startswith(("est", "derived")) for e in inp.earnings)
    periods = {e.period: e for e in inp.earnings}
    assert periods["FY2014"].eps_cents == 260.0 and periods["FY2016"].eps_cents == -120.0
    assert "1H2016 TTM" not in periods                                     # excluded, with reason
    assert periods["1H2017 TTM"].eps_cents == pytest.approx(-120.0 + 106.5 + 60.2, abs=0.06)   # SWS 46.749 vs BHP figures to 0.1c
    assert periods["FY2026"].date == D(2026, 8, 18)
    assert any("left out" in w for w in inp.warnings) and any("FY2017-FY2025" in w for w in inp.warnings)


def test_sws_import_half_year_anchor_and_deadline():
    from src.sws_import import annual_eps_rows, asx_deadline
    assert asx_deadline(D(2025, 12, 31)) == D(2026, 2, 28)
    assert asx_deadline(D(2023, 12, 31)) == D(2024, 2, 29)
    assert asx_deadline(D(2026, 6, 30)) == D(2026, 8, 31)
    rows = {
        "past_last_earnings_update": [1767139200],          # 2025-12-31 (a half-year end)
        "past_last_earnings_update_annual": [1751241600],   # 2025-06-30
        "past_earnings_per_share_ltm_history_offset": [0, -1],
        "past_earnings_per_share_ltm_history": [2.01813, 2.24554],
        "dividend_historical_dividend_payments_date": [1771286400, 1744243200],  # 2026-02-17, 2025-04-10
        "currency_info_reporting_currency_iso": ["USD"],
    }
    out = annual_eps_rows(rows, "SWS")
    assert [(r["period"], r["date"]) for r in out] == [("1H2025 TTM", "2025-02-28"), ("1H2026 TTM", "2026-02-17")]
    assert out[1]["ttm_eps"] == pytest.approx(201.813) and out[1]["currency"] == "USD"


def test_smoothed_eps_uses_only_what_was_known_and_skips_one_offs():
    from src.value_wedge import smooth_eps
    e = [EarningsPoint(D(2012, 8, 31), "FY2012", 300.0), EarningsPoint(D(2013, 8, 31), "FY2013", 200.0),
         EarningsPoint(D(2014, 8, 31), "FY2014", 100.0), EarningsPoint(D(2015, 8, 31), "FY2015", -120.0),
         EarningsPoint(D(2016, 8, 31), "FY2016", 160.0)]
    out = {x.period: x for x in smooth_eps(e, years=3, exclude=["FY2015"])}
    assert out["FY2012"].norm_eps_cents == 300.0                      # only itself known
    assert out["FY2014"].norm_eps_cents == pytest.approx(200.0)       # 300, 200, 100
    assert out["FY2015"].norm_eps_cents == pytest.approx(200.0)       # loss left out: 300, 200, 100
    assert out["FY2016"].norm_eps_cents == pytest.approx(130.0)       # 100, 160 (FY2013 aged out)
    assert out["FY2015"].eps_used == pytest.approx(200.0) and out["FY2015"].eps_cents == -120.0


def test_fair_line_sits_30pct_up_the_wedge():
    e = [EarningsPoint(D(2020, 1, 1), "FY19", 100.0)]
    m = [MultipleRegime(D(2020, 1, 1), 13, 20)]
    r = compute_wedge(_inputs(e, m, _prices(value=15.0)))
    cur = r.current
    assert cur.fair_pe == pytest.approx(15.1) and cur.fair_value == pytest.approx(15.1)
    assert "fair $15.10 (15.1x)" in r.headline


def test_bhp_smoothed_corridor_with_statutory_kept():
    from src.value_wedge import daily_corridor
    px = _prices("2013-08-01", "2026-09-30", value=45.0)
    fx = pd.Series(1.5, index=px.index)
    inp = load_wedge_inputs("BHP", repo=repo_root(), prices=px, fx=fx, start="year_before_buy")
    assert inp.smoothing and "FY2016" in inp.smoothing
    fy16 = next(e for e in inp.earnings if e.period == "FY2016")
    assert fy16.eps_cents == -120.0 and fy16.norm_eps_cents > 0      # statutory loss, smoothed still positive
    day = daily_corridor(compute_wedge(inp))
    d = day.loc["2016-10-03"]
    assert d["buy"] > 0 and (d["buy_reported"] != d["buy_reported"])  # smoothed corridor on; statutory NaN (loss)
    assert {"fair", "fair_pe", "eps_reported", "buy_reported", "sell_reported"} <= set(day.columns)


def test_forecast_extends_corridor_right_of_today_without_touching_history(tmp_path):
    from src.value_wedge import forecast_corridor, forecast_lines
    from src.wedge_chart import render_wedge_chart
    from src.wedge_report import wedge_markdown
    from src.wedge_workbook import build_wedge_workbook
    e = [EarningsPoint(D(2020, 1, 1), "FY2019", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 50)]
    inp = _inputs(e, m, _prices(value=10.0))
    inp.fair_split = 0.35
    inp.forecast = [{"period": "FY2024", "date": D(2024, 8, 20), "eps_cents": 30.0, "note": "margin back"},
                    {"period": "FY2025", "date": D(2025, 8, 20), "eps_cents": 40.0, "note": ""}]
    inp.forecast_label = "Recovery case"
    r = compute_wedge(inp)
    # history untouched: today's corridor is still on reported EPS
    assert r.current.buy_value == pytest.approx(5.0) and r.current.zone == "Upper third of wedge"
    f24, f25 = r.forecast
    assert (f24["buy_value"], f24["sell_value"]) == (pytest.approx(7.5), pytest.approx(15.0))
    assert f24["fair_value"] == pytest.approx(30.0 * (25 + 0.35 * 25) / 100)
    assert f25["forward_pe_today"] == pytest.approx(25.0) and f25["zone_today"] == "Lower third of wedge"
    fc = forecast_corridor(r)
    assert fc.index[0] > pd.Timestamp(r.current.date)
    assert fc.loc[:"2024-08-19", "buy"].eq(5.0).all()            # today's EPS until the first release
    assert fc.loc["2024-08-20":"2025-08-19", "buy"].eq(7.5).all()
    assert fc["buy"].iloc[-1] == pytest.approx(10.0)
    assert "FY2025F" in forecast_lines(r)[1] and "not reported" in wedge_markdown(r)
    assert "FY2024F" in wedge_html(r)
    render_wedge_chart(r, tmp_path / "w.png")
    svg = (tmp_path / "w.svg").read_text()
    assert "FY25F" in svg and "today | forecast" in svg
    from openpyxl import load_workbook
    wb = load_workbook(build_wedge_workbook(r, tmp_path / "w.xlsx"))
    assert wb["Forecast"]["H2"].value == "=D2*E2/100"


def test_fy_labels_sit_under_the_price_panel_too(tmp_path):
    from src.wedge_chart import render_wedge_chart
    e = [EarningsPoint(D(2020, 1, 1), "FY2019", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 50)]
    render_wedge_chart(compute_wedge(_inputs(e, m, _prices())), tmp_path / "w.png")
    svg = (tmp_path / "w.svg").read_text()
    assert svg.count(">FY22<") == 2          # under the price chart and under the PE panel


def test_pwh_forecast_is_labelled_and_not_reported():
    inp = load_wedge_inputs("PWH", repo=repo_root(), prices=_prices("2016-01-04", "2026-10-01"))
    if not inp.forecast:
        pytest.skip("no forecast block in pwh_ax.yaml")
    assert all(f["date"] > D(2026, 10, 1) for f in inp.forecast)
    assert any("Not reported" in w for w in inp.warnings)


def test_closed_ledger_decisions_mark_the_sale(tmp_path):
    import json
    from src.value_wedge import decision_notes
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "decisions.json").write_text(json.dumps({"decisions": [
        {"ticker": "TST", "date": "2024-01-13", "status": "CLOSED",
         "flows": [["2024-01-13", -1.0], ["2024-04-19", 0.02], ["2026-05-06", 0.56]]},
        {"ticker": "TST", "date": "2024-06-06", "status": "CLOSED",
         "flows": [["2024-06-06", -1.0], ["2026-05-06", 0.7]]},
        {"ticker": "TST", "date": "2025-01-01", "status": "OPEN", "flows": [["2025-01-01", -1.0]]}]}))
    notes = decision_notes(tmp_path, "TST.AX")
    sold = [n for n in notes if n.label.startswith("Sold")]
    assert len(sold) == 1 and sold[0].date == D(2026, 5, 6) and sold[0].label == "Sold (decisions 1, 2)"
    assert [n.label for n in notes].count("Bought (decision 3 of 3)") == 1


def test_lau_config_is_eodhd_statutory():
    inp = load_wedge_inputs("LAU", repo=repo_root(), prices=_prices("2016-01-04", "2026-10-09"))
    fy26 = next(e for e in inp.earnings if e.period == "FY2026")
    assert fy26.eps_cents == pytest.approx(6.245) and "EODHD" in fy26.source
    assert any(n.label.startswith("Sold") and n.date == D(2026, 5, 6) for n in inp.notes)


def test_low_line_names_the_longest_multi_year_low():
    from src.value_wedge import low_line
    idx = pd.bdate_range("2014-01-01", "2026-10-09")
    s = pd.Series(30.0, index=idx)
    s[s.index < "2017-01-01"] = 12.0          # lower 10 years ago only
    s.iloc[-1] = 16.9
    e = [EarningsPoint(D(2014, 1, 1), "FY2013", 100.0)]
    m = [MultipleRegime(D(2014, 1, 1), 17, 30)]
    r = compute_wedge(_inputs(e, m, s))
    assert r.price_low["years"] == 5 and "At its 5-year low" in low_line(r) and "Dec 2016" in low_line(r)


def test_forecast_cases_on_a_smoothed_basis_and_outlook_in_the_email(tmp_path):
    import yaml
    from src.wedge_report import wedge_html, wedge_markdown
    from src.wedge_chart import render_wedge_chart
    from src.value_wedge import forecast_corridor
    (tmp_path / "valuations").mkdir()
    (tmp_path / "valuations" / "tst_ax.yaml").write_text(yaml.safe_dump({
        "ticker": "TST.AX", "company_name": "Test", "buy_multiple": "archetype", "sell_multiple": 30,
        "earnings": [{"date": f"{y}-08-20", "period": f"FY{y}", "period_end": f"{y}-06-30", "ttm_eps": v}
                     for y, v in ((2021, 120.0), (2022, 150.0), (2023, 108.0), (2024, 125.0), (2025, 117.0),
                                  (2026, 111.0))],
        "forecast": {"extend_to": "2029-02-28", "cases": [
            {"label": "Recovery", "basis": "consensus", "rows": [
                {"period": "FY2027", "release": "2027-08-20", "eps": 113.0},
                {"period": "FY2028", "release": "2028-08-20", "eps": 125.0}]},
            {"label": "No recovery", "basis": "trend", "rows": [
                {"period": "FY2027", "release": "2027-08-20", "eps": 104.0},
                {"period": "FY2028", "release": "2028-08-20", "eps": 98.0}]}]},
        "outlook": {"as_of": "12 Oct 2026", "recover": ["OEM supply back"], "not_recover": ["Diesel utes shrinking"]},
    }))

    class Cls:   # what shared.classifier would return
        buy_multiple, archetype, eps_basis, base_multiple = 17.0, "defensive_mature", "avg_3y", 15
        modifiers, guardrails, suggested, error = [], [], True, None

    inp = load_wedge_inputs("TST", repo=tmp_path, prices=_prices("2021-01-04", "2026-10-09", value=16.9),
                            classification=Cls())
    r = compute_wedge(inp)
    assert r.current.buy_pe == 17 and r.current.eps_used == pytest.approx((125 + 117 + 111) / 3, abs=0.01)
    rec = {(f["case"], f["period"]): f for f in r.forecast}
    assert rec[("Recovery", "FY2028")]["eps_basis_cents"] == pytest.approx((111 + 113 + 125) / 3, abs=0.01)
    assert rec[("No recovery", "FY2028")]["buy_value"] == pytest.approx((111 + 104 + 98) / 3 * 17 / 100, abs=0.01)
    assert not forecast_corridor(r, "No recovery").empty
    html = wedge_html(r)
    assert "Why it might recover" in html and "Diesel utes shrinking" in html and ">No recovery<" in html
    for banned in ("display:flex", "display:grid", "<style", "class="):
        assert banned not in html
    assert "**Why it might not**" in wedge_markdown(r)
    render_wedge_chart(r, tmp_path / "w.png")
    assert "No recovery FY28F" in (tmp_path / "w.svg").read_text()


class _Cls:
    """What shared.classifier returns, for offline tests."""
    def __init__(self, bm=17.0, arch="quality_steady", basis="ttm", median=28.4, alt=None):
        self.buy_multiple, self.archetype, self.eps_basis, self.base_multiple = bm, arch, basis, 15
        self.modifiers, self.guardrails, self.suggested, self.error = [], [], True, None
        self.median_pe_10y, self.alt_metric = median, alt


def _eodhd_rows(vals, ccy="AUD"):
    rows = [{"date": f"{y}-08-20", "period": f"FY{y}", "period_end": f"{y}-06-30", "ttm_eps": v,
             "ttm_div": None, "source": "EODHD income statement", "notes": ""} for y, v in vals]
    return rows, ccy


def test_holding_without_a_file_gets_an_automatic_config(tmp_path, monkeypatch):
    import shared.eps_history as eh
    from shared import eodhd
    monkeypatch.setattr(eh, "eodhd_earnings_rows",
                        lambda t, **k: _eodhd_rows([(2023, 30.0), (2024, 34.0), (2025, 36.0), (2026, 40.0)]))
    trends = [{"date": "2027-06-30", "period": "0y", "earningsEstimateAvg": "0.44",
               "earningsEstimateNumberOfAnalysts": "9", "earningsEstimateLow": "0.40", "earningsEstimateHigh": "0.48"},
              {"date": "2027-06-30", "period": "+1y", "earningsEstimateAvg": "0.39",
               "earningsEstimateNumberOfAnalysts": "8"},
              {"date": "2028-06-30", "period": "+1y", "earningsEstimateAvg": "0.50",
               "earningsEstimateNumberOfAnalysts": "8"},
              {"date": "2026-06-30", "period": "0y", "earningsEstimateAvg": "0.41",
               "earningsEstimateNumberOfAnalysts": "9"}]
    inp = load_wedge_inputs("TST", repo=tmp_path, prices=_prices("2022-01-03", "2026-10-09", value=8.0),
                            classification=_Cls(), fund={"General": {"Name": "Test Holdings"}}, trend_rows=trends)
    r = compute_wedge(inp)
    assert inp.company_name == "Test Holdings"
    assert r.current.buy_pe == 17 and r.current.sell_pe == 28.4          # SELL = own 10-year median
    assert any("Automatic config" in w for w in inp.warnings)
    fy27 = next(f for f in r.forecast if f["period"] == "FY2027")
    assert fy27["eps_cents"] == pytest.approx(44.0) and "9 analysts" in fy27["note"]   # "0y" beats the old "+1y"
    assert [f["period"] for f in r.forecast] == ["FY2027", "FY2028"]      # the reported FY26 estimate is dropped


def test_pre_profit_and_loss_making_holdings_say_why(tmp_path, monkeypatch):
    import shared.eps_history as eh
    from src.value_wedge import WedgeConfigMissing
    monkeypatch.setattr(eh, "eodhd_earnings_rows", lambda t, **k: _eodhd_rows([(2025, 3.0), (2026, -2.0)]))
    with pytest.raises(WedgeConfigMissing, match="loss-making"):
        load_wedge_inputs("TST", repo=tmp_path, prices=_prices(), classification=_Cls(), fund={}, trend_rows=[])
    with pytest.raises(WedgeConfigMissing, match="no P/E line"):
        load_wedge_inputs("TST", repo=tmp_path, prices=_prices(), classification=_Cls(bm=None, alt="price_to_sales"),
                          fund={}, trend_rows=[])


def test_sell_rule_keeps_sell_between_1_5_and_3_times_buy():
    from shared.classifier import sell_multiple_for
    assert sell_multiple_for(_Cls(bm=17, median=28.4))[0] == 28.4
    assert sell_multiple_for(_Cls(bm=10, median=8))[0] == 15.0           # floor 1.5x BUY
    assert sell_multiple_for(_Cls(bm=10, median=45))[0] == 30.0          # cap 3x BUY
    assert sell_multiple_for(_Cls(bm=10, median=None))[0] == 15.0


def test_earnings_source_eodhd_replaces_the_hand_typed_rows(tmp_path, monkeypatch):
    import yaml
    import shared.eps_history as eh
    monkeypatch.setattr(eh, "eodhd_earnings_rows", lambda t, **k: _eodhd_rows([(2024, 190.0), (2025, 200.0)], "USD"))
    (tmp_path / "valuations").mkdir()
    (tmp_path / "valuations" / "tst_ax.yaml").write_text(yaml.safe_dump({
        "ticker": "TST.AX", "buy_multiple": 9, "sell_multiple": 16, "earnings_source": "eodhd", "auto_forecast": False,
        "earnings": [{"date": "2024-02-01", "period": "FY2023", "ttm_eps": 999, "notes": "Est."}]}))
    fx = pd.Series(1.5, index=pd.bdate_range("2023-01-02", "2026-10-09"))
    inp = load_wedge_inputs("TST", repo=tmp_path, prices=_prices("2023-01-02", "2026-10-09"), fx=fx)
    assert [e.eps_cents for e in inp.earnings] == [190.0, 200.0] and inp.eps_currency == "USD"
    assert not any(e.eps_cents == 999 for e in inp.earnings)


def test_one_pdf_for_the_run(tmp_path):
    from src.wedge_chart import render_wedge_chart
    from src.wedge_pdf import build_report_pdf
    e = [EarningsPoint(D(2020, 1, 1), "FY2019", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 50)]
    folder = tmp_path / "TST"
    render_wedge_chart(compute_wedge(_inputs(e, m, _prices())), folder / "wedge_chart.png")
    (folder / "memo.md").write_text("# Review\n\n## 1. Why\n- near the high\n\n" + "Long paragraph. " * 400)
    rows = [{"ticker": "TST", "alert_tier": "Tier 1: Watch", "distance_to_high_pct": 3.1, "wedge_zone": "Above SELL",
             "wedge_position": 1.2, "wedge_buy_value": 5.0, "wedge_fair_value": 6.0, "wedge_sell_value": 10.0,
             "sally_verdict": "Watch only"},
            {"ticker": "PRE", "alert_tier": "Tier 1: Watch", "distance_to_high_pct": 1.0, "wedge_zone": None,
             "wedge_position": None, "wedge_missing": "no wedge: pre-profit", "sally_verdict": "Watch only"}]
    out = build_report_pdf(rows, tmp_path, tmp_path / "r.pdf", "Sally", "2 companies")
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and data.count(b"/Type /Page") >= 4    # cover + chart + memo pages


def test_workbook_chart_looks_like_the_wedge(tmp_path):
    from openpyxl import load_workbook
    from src.wedge_workbook import build_wedge_workbook
    e = [EarningsPoint(D(2020, 1, 1), "FY2019", 20.0)]
    m = [MultipleRegime(D(2020, 1, 1), 25, 50)]
    inp = _inputs(e, m, _prices(value=4.0))
    inp.forecast = [{"case": "Consensus", "period": "FY2024", "date": D(2024, 8, 20), "eps_cents": 30.0, "note": ""}]
    inp.forecast_cases = [{"label": "Consensus", "basis": ""}]
    out = build_wedge_workbook(compute_wedge(inp), tmp_path / "w.xlsx")
    wb = load_workbook(out)
    ws = wb["Weekly"]
    hdr = [c.value for c in ws[1]]
    assert "Chart: buying zone" in hdr and "Chart: forecast band" in hdr and "Fair value" in hdr
    fc_rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 23).value == 1]
    assert fc_rows and ws.cell(fc_rows[0], 2).value is None             # forecast weeks: no price
    assert "Forecast!" in ws.cell(fc_rows[-1], 3).value                  # EPS from the Forecast sheet
    import zipfile
    xml = zipfile.ZipFile(out).read("xl/charts/chart1.xml").decode()
    assert "areaChart" in xml and 'grouping val="stacked"' in xml and "lineChart" in xml and "FY" in xml
